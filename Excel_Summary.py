# -*- coding: utf-8 -*-
"""
Друга програма: підсумковий блок по категоріях (з рядка 143 і нижче).

Бере файл, який вийшов після першої програми (колонка C — номер категорії, D — кількість штук),
і додає в кінець блок «№ Категорії / Кількість / формула / Сума»:
  * блок починається на 4 рядки нижче рядка «Разом» (Разом у рядку 139 -> блок з 143)
    і не заходить на нижчі рядки, якщо в них уже щось внесено вручну;
  * A–E (назва, одиниця, код, номер категорії) — постійні, переносяться як є;
  * F — сума колонки D по категорії (формула SUMIF), I — сума грошей (колонка «Сальдо на кінець»,
    Дебет) по категорії; для «послуг» (6) кількості немає, є тільки сума;
  * G — формула: кількість/1000/2*0,175 (т), кількість/2*0,175 (кг) або Сума/1000 (грн.);
  * нижче — підсумки, кольори та решта комірок зі зразка (1-ПЕ тощо).
Оригінал не змінюється: поруч зберігається «підсумок_<назва>».

Запуск: перетягнути файл на програму (Excel_Summary.exe або Drop_file_summary.bat) або двічі клікнути.
Потрібні пакети: pip install openpyxl xlrd xlutils xlwt   (і файл Excel_Accounting.py поруч)
"""
from __future__ import annotations

import datetime as dt
import os
import re
import sys
from decimal import Decimal
from pathlib import Path

from Excel_Accounting import (CAT_COL, MARK, TARGET_COL, TOTAL_WORDS, UserError, XlsSheet, XlsxSheet,
                              find_columns, pick_files, say, to_decimal)

OUT_PREFIX = "підсумок_"
BLOCK_ROW0 = 143                 # рядок шаблону, з якого починається блок
GAP_AFTER_TOTAL = 4              # «Разом» 139 -> блок з 143
LABEL_HEADER = "№ Категорії"
MONTHS = {"січень": 1, "лютий": 2, "березень": 3, "квітень": 4, "травень": 5, "червень": 6, "липень": 7,
          "серпень": 8, "вересень": 9, "жовтень": 10, "листопад": 11, "грудень": 12}
LET = "ABCDEFGHIJKL"

# ───────────────────────── вбудований шаблон блоку (зі зразка + кольори зі скріншота) ─────────────────────────
# STYLES: шрифт [ім'я, висота у твіпах, жирний, курсив, RGB], заливка RGB, рамки [стиль, RGB] (верх, низ, ліво, право),
# вирівнювання [гориз., верт., перенос], формат числа. CELLS: (зсув рядка, колонка, номер стилю).
STYLES = [{"font":["Arial",280,1,0,None],"fill":None,"bord":[[5,None],[5,None],[5,None],[1,None]],"al":[1,2,0],"fmt":"mmm-yy"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[5,None],[5,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Times New Roman",180,0,0,None],"fill":"C6E0B4","bord":[[5,None],[5,None],[1,None],[1,None]],"al":[2,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[5,None],[5,None],[1,None],[5,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[1,None],[5,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[1,None],[1,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"F8CBAD","bord":[[0,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[0,None],[1,None],[1,None],[1,None]],"al":[2,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[1,None],[1,None],[5,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[1,None],[5,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"F8CBAD","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[1,None],[1,None],[1,None],[1,None]],"al":[2,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[1,None],[1,None],[5,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[5,None],[0,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"D9E1F2","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"D9E1F2","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":"D9E1F2","bord":[[1,None],[1,None],[1,None],[5,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"DBDBDB","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",200,0,0,None],"fill":"DBDBDB","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":"DBDBDB","bord":[[1,None],[1,None],[1,None],[5,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[1,None],[1,None],[5,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"A9D08E","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",220,0,0,[255,0,0]],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[1,None],[5,None],[5,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[1,None],[5,None],[1,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[5,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"A9D08E","bord":[[1,None],[0,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[0,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[0,None],[1,None],[5,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",200,0,0,None],"fill":"A9D08E","bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":"F8CBAD","bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[5,None],[5,None],[1,None],[1,None]],"al":[2,2,0],"fmt":"General"},{"font":["Arial",200,1,0,None],"fill":None,"bord":[[5,None],[0,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[5,None],[0,None],[1,None],[0,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[5,None],[0,None],[5,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[5,None],[0,None],[1,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"A9D08E","bord":[[5,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[5,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":"FFFF00","bord":[[5,None],[1,None],[1,None],[1,None]],"al":[2,2,0],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[5,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[5,None],[0,None],[1,None],[5,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",160,1,0,None],"fill":None,"bord":[[0,None],[0,None],[5,None],[0,None]],"al":[1,1,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,1,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[1,None],[1,None],[0,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[1,None],[1,None],[0,None],[1,None]],"al":[1,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[1,None],[1,None],[1,None]],"al":[0,2,1],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":"FFFF00","bord":[[0,None],[1,None],[1,None],[1,None]],"al":[2,2,0],"fmt":"General"},{"font":["Arial",160,1,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,1,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[1,None],[1,None],[0,None],[0,None]],"al":[0,2,1],"fmt":"General"},{"font":["Arial",200,0,0,None],"fill":"D9D9D9","bord":[[1,None],[1,None],[0,None],[1,None]],"al":[0,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":"FFFF00","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[2,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",180,1,0,None],"fill":"FFFF00","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",180,0,0,None],"fill":None,"bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",180,0,0,None],"fill":None,"bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.0"},{"font":["Arial",180,1,0,None],"fill":"FFFF00","bord":[[1,None],[1,None],[1,None],[1,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",280,1,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",160,0,0,[255,0,0]],"fill":"FFFF00","bord":[[0,None],[0,None],[0,None],[0,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",160,1,0,None],"fill":None,"bord":[[0,None],[1,None],[0,None],[0,None]],"al":[2,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[0,None],[1,None],[0,None],[0,None]],"al":[1,2,0],"fmt":"General"},{"font":["Arial",160,0,0,[255,0,0]],"fill":None,"bord":[[0,None],[1,None],[0,None],[0,None]],"al":[1,2,0],"fmt":"0.00"},{"font":["Arial",200,0,0,[0,128,128]],"fill":None,"bord":[[1,[192,192,192]],[1,[192,192,192]],[1,[192,192,192]],[1,[192,192,192]]],"al":[3,0,1],"fmt":"#,##0.00"},{"font":["Arial",180,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.00"},{"font":["Arial",200,0,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.00"},{"font":["Arial",180,0,0,[255,0,0]],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.00"},{"font":["Arial",180,0,0,None],"fill":"FFFF00","bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.00"},{"font":["Arial",160,0,0,[153,51,0]],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"General"},{"font":["Arial",180,1,0,None],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.0"},{"font":["Arial",160,0,0,[153,51,0]],"fill":None,"bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.0"},{"font":["Arial",240,1,0,None],"fill":"FFFF00","bord":[[0,None],[0,None],[0,None],[0,None]],"al":[0,2,0],"fmt":"0.0"}]
CELLS = [[0,0,0],[0,1,1],[0,2,1],[0,3,1],[0,4,1],[0,5,1],[0,6,2],[0,7,1],[0,8,3],[1,0,6],[1,1,7],[1,2,8],[1,3,9],[1,4,10],[1,5,8],[1,6,9],[1,7,11],[1,8,12],[2,0,13],[2,1,14],[2,2,15],[2,3,16],[2,4,17],[2,5,15],[2,6,16],[2,7,18],[2,8,19],[2,9,20],[3,0,13],[3,1,14],[3,2,15],[3,3,16],[3,4,17],[3,5,15],[3,6,22],[3,7,23],[3,8,24],[3,9,4],[3,10,25],[4,0,13],[4,1,14],[4,2,15],[4,3,16],[4,4,17],[4,5,15],[4,6,16],[4,7,18],[4,8,19],[5,0,13],[5,1,14],[5,2,15],[5,3,16],[5,4,17],[5,5,15],[5,6,16],[5,7,18],[5,8,19],[6,0,13],[6,1,14],[6,2,16],[6,3,16],[6,4,17],[6,5,26],[6,6,27],[6,7,28],[6,8,29],[7,0,30],[7,1,31],[7,2,32],[7,3,16],[7,4,17],[7,5,32],[7,6,16],[7,7,18],[7,8,19],[8,0,30],[8,1,31],[8,2,32],[8,3,16],[8,4,17],[8,5,32],[8,6,16],[8,7,18],[8,8,19],[9,0,30],[9,1,31],[9,2,32],[9,3,16],[9,4,17],[9,5,32],[9,6,16],[9,7,18],[9,8,19],[10,0,30],[10,1,31],[10,2,32],[10,3,16],[10,4,17],[10,5,32],[10,6,16],[10,7,18],[10,8,19],[11,0,34],[11,1,35],[11,2,32],[11,3,36],[11,4,17],[11,5,37],[11,6,16],[11,7,38],[11,8,39],[11,10,40],[11,11,41],[12,1,43],[12,2,44],[12,3,1],[12,4,45],[12,5,1],[12,6,1],[12,7,46],[12,8,47],[13,0,49],[13,1,50],[13,2,51],[13,3,52],[13,4,53],[13,5,51],[13,6,52],[13,7,54],[13,8,55],[13,9,56],[14,0,58],[14,1,59],[14,2,8],[14,3,60],[14,4,61],[14,5,8],[14,6,9],[14,7,11],[14,8,19],[14,10,63],[14,11,5],[15,0,64],[15,1,65],[15,2,32],[15,3,21],[15,4,66],[15,5,32],[15,6,16],[15,7,18],[15,8,19],[16,0,14],[16,1,14],[16,2,16],[16,3,16],[16,4,16],[16,5,16],[16,6,16],[16,7,16],[16,8,67],[16,10,40],[16,11,41],[17,0,16],[17,1,16],[17,2,16],[17,3,16],[17,4,16],[17,5,68],[17,6,69],[17,7,70],[17,8,71],[19,5,73],[21,4,75],[21,5,76],[21,6,76],[21,7,76],[22,2,77],[22,3,78],[22,4,75],[22,5,79],[22,6,79],[22,7,79],[22,8,80],[24,5,83],[24,6,80],[24,7,83],[24,8,80],[24,10,81],[26,8,85],[26,9,86],[26,10,87]]
HEIGHTS = {"0":525,"1":780,"2":510,"3":390,"4":750,"5":750,"6":390,"7":765,"8":390,"9":795,"10":510,"11":780,"12":285,"13":540,"14":330,"15":525,"16":255,"17":255,"18":255,"19":360,"20":255,"21":255,"22":255,"23":255,"24":255,"25":255,"26":315}
MERGES = [[20,0,1],[21,0,1],[22,0,1],[11,0,1],[13,0,1],[14,0,1],[15,0,1],[16,0,1],[6,0,1],[7,0,1],[8,0,1],[9,0,1],[10,0,1],[1,0,1],[2,0,1],[3,0,1],[4,0,1],[5,0,1]]
CAT_ROWS = [[1,"Проспекти, плакати та інша рекламна продукція, віддрукована (крім каталогів рекламних)","т","18.12.12.50.A",1],[2,"Календарі усіх видів, уключаючи блокові календарі, віддруковані","т","18.12.19.10.A",2],[3,"Продукція інша, н.в.і.у., віддрукована","грн.","18.12.19.90",3],[4,"Книги, брошури, листівки та подібна продукція, у вигляді окремих аркушів, віддруковані","т","18.12.14.07.A",4],[5,"Книги, брошури, листівки та подібна продукція, віддруковані (крім у вигляді окремих аркушів)","т","18.12.14.14.A",5],[6,"Послуги поліграфічні інші","грн.","18.13.30.00",6],[7,"Вироби пакувальні інші з паперу та картону н.в.і.у., включаючи конверти для грамплатівок","кг","17.21.15.30",7],[8,"Конверти з паперу чи картону","кг","17.23.12.30",8],[9,"Швидкозшивачі (оправи знімні), папки, обкладинки з паперу чи картону (крім обкладинок для книг)","кг","17.23.13.50",9],[10,"Етикетки та ярлики з паперу чи картону друковані, самосклеювальні","кг","17.29.11.20",10],[11,"Етикетки та ярлики з паперу чи картону друковані інші (крім самосклеювальних)","кг","17.29.11.40",11],[13,"Підноси, тарелі, тарілки, чашки та вироби подібні, з паперу або картону","кг","17.22.13.00",12],[14,"Каталоги рекламні, віддруковані","т","18.12.12.30.A",13],[15,"Книжки записні, блокноти для листів та записів, з паперу чи картону","кг","17.23.13.15",14]]
WIDTHS = {"7":2486,"8":3035,"10":2925}

# ───────────────────────── запис блоку в .xls / .xlsx ─────────────────────────
def _hex(rgb):
    return "%02X%02X%02X" % tuple(rgb)


class XlsBlock(XlsSheet):
    """.xls: стилі з точними кольорами (вільні комірки палітри перевизначаються)."""

    def __init__(self, path):
        super().__init__(path)
        import xlwt
        self.xlwt = xlwt
        self._styles = {}
        used = set()
        for xf in self.rb.xf_list:
            b, bd = xf.background, xf.border
            used.update((b.pattern_colour_index, b.background_colour_index, bd.top_colour_index, bd.bottom_colour_index,
                         bd.left_colour_index, bd.right_colour_index))
        used.update(f.colour_index for f in self.rb.font_list)
        self._free = [i for i in range(8, 64) if i not in used]
        self._cmap = {i: tuple(self.rb.colour_map[i]) for i in range(64) if self.rb.colour_map.get(i)}
        self._have = {}

    def color(self, rgb):
        rgb = tuple(rgb)
        if rgb in self._have:
            return self._have[rgb]
        for i, v in sorted(self._cmap.items()):
            if v == rgb:
                self._have[rgb] = i
                return i
        if self._free:
            i = self._free.pop(0)
            self.wb.set_colour_RGB(i, *rgb)
            self._cmap[i] = rgb
        else:                                        # вільних немає — найближчий колір палітри
            i = min((k for k in self._cmap if k >= 1), key=lambda k: sum((a - b) ** 2 for a, b in zip(self._cmap[k], rgb)))
        self._have[rgb] = i
        return i

    def style(self, sid):
        if sid in self._styles:
            return self._styles[sid]
        x = self.xlwt
        s = STYLES[sid]
        st = x.XFStyle()
        name, height, bold, italic, frgb = s["font"]
        st.font.name, st.font.height, st.font.bold, st.font.italic = name, height, bool(bold), bool(italic)
        if frgb:
            st.font.colour_index = self.color(frgb)
        st.num_format_str = s["fmt"]
        st.alignment.horz, st.alignment.vert, st.alignment.wrap = s["al"][0], s["al"][1], s["al"][2]
        for side, (ls, rgb) in zip(("top", "bottom", "left", "right"), s["bord"]):
            setattr(st.borders, side, ls)
            if ls and rgb:
                setattr(st.borders, side + "_colour", self.color(rgb))
        if s["fill"]:
            st.pattern.pattern = x.Pattern.SOLID_PATTERN
            st.pattern.pattern_fore_colour = self.color(bytes.fromhex(s["fill"]))
        self._styles[sid] = st
        return st

    def put_block(self, r, c, value, sid):
        v = value
        if isinstance(value, tuple):                 # ("f", "формула")
            v = self.xlwt.Formula(value[1])
        self.ws.write(r - 1, c - 1, v, self.style(sid))
        if isinstance(value, tuple):                 # перерахунок при відкритті
            self.ws._Worksheet__rows[r - 1]._Row__cells[c - 1].calc_flags = 3

    def row_height(self, r, twips):
        row = self.ws.row(r - 1)
        row.height_mismatch = True
        row.height = twips

    def col_width(self, c, width):
        info = self.rs.colinfo_map.get(c - 1)
        if info is None or info.width < width:
            self.ws.col(c - 1).width = width

    def merge(self, r, c1, c2):
        self.ws.merged_ranges.append((r - 1, r - 1, c1 - 1, c2 - 1))


class XlsxBlock(XlsxSheet):
    def put_block(self, r, c, value, sid):
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        s = STYLES[sid]
        cell = self.ws.cell(r, c)
        name, height, bold, italic, frgb = s["font"]
        cell.font = Font(name=name, size=height / 20, bold=bool(bold), italic=bool(italic), color=_hex(frgb) if frgb else None)
        if s["fill"]:
            cell.fill = PatternFill("solid", fgColor=s["fill"])
        kinds = {0: None, 1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick", 6: "double", 7: "hair"}
        sides = {k: Side(style=kinds.get(ls), color=_hex(rgb) if rgb else None)
                 for k, (ls, rgb) in zip(("top", "bottom", "left", "right"), s["bord"])}
        cell.border = Border(**sides)
        hor = {0: "general", 1: "left", 2: "center", 3: "right"}.get(s["al"][0], "general")
        ver = {0: "top", 1: "center", 2: "bottom"}.get(s["al"][1], "bottom")
        cell.alignment = Alignment(horizontal=hor, vertical=ver, wrap_text=bool(s["al"][2]))
        cell.number_format = s["fmt"]
        if value is not None:
            cell.value = ("=" + value[1]) if isinstance(value, tuple) else value

    def row_height(self, r, twips):
        self.ws.row_dimensions[r].height = twips / 20

    def col_width(self, c, width):
        from openpyxl.utils import get_column_letter
        d = self.ws.column_dimensions[get_column_letter(c)]
        if not d.width or d.width < width / 256:
            d.width = width / 256

    def merge(self, r, c1, c2):
        self.ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)


def open_block_sheet(path: Path):
    ext = path.suffix.lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            return XlsxBlock(path)
        if ext == ".xls":
            return XlsBlock(path)
    except ImportError as e:
        raise UserError(f"Не вистачає пакета: {e.name}.\nВстановіть:  pip install openpyxl xlrd xlutils xlwt")
    raise UserError(f"Формат «{ext}» не підтримується. Потрібен .xls або .xlsx.")


# ───────────────────────── допоміжне ─────────────────────────
def find_money_col(sh, head: int) -> int:
    """Колонка «Сальдо на кінець періоду → Дебет» (гроші)."""
    for r in range(1, head + 1):
        for c in range(1, 15):
            v = sh.get(r, c)
            if isinstance(v, str) and v.strip().casefold().startswith("сальдо на кінець"):
                for cc in (c, c + 1):
                    d = sh.get(head, cc)
                    if isinstance(d, str) and d.strip().casefold() == "дебет":
                        return cc
    raise UserError("Не знайшов колонку «Сальдо на кінець періоду → Дебет». Скрипт не вгадує колонки.")


def report_month(sh):
    for r in range(1, 8):
        for c in range(1, 4):
            v = sh.get(r, c)
            if isinstance(v, str):
                m = re.search(r"за\s+([A-Za-zА-Яа-яІіЇїЄєҐґ]+)\s+(\d{4})", v)
                if m and m.group(1).casefold() in MONTHS:
                    return dt.date(int(m.group(2)), MONTHS[m.group(1).casefold()], 1)
    return None


def last_filled_row(sh) -> int:
    for r in range(sh.nrows, 0, -1):
        if any(sh.get(r, c) not in (None, "") for c in range(1, 16)):
            return r
    return 0


def num(v):
    return Decimal(str(v)) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


# ───────────────────────── основна обробка ─────────────────────────
def process(src: Path):
    sh = open_block_sheet(src)
    name_col, _, head = find_columns(sh)
    money_col = find_money_col(sh, head)

    total_row = None
    for r in range(head + 1, sh.nrows + 1):
        v = sh.get(r, name_col)
        if isinstance(v, str) and v.strip().casefold().strip(" :.") in TOTAL_WORDS:
            total_row = r
            break
    last = last_filled_row(sh)
    for r in range(head + 1, last + 1):
        if sh.get(r, 5) == LABEL_HEADER:
            raise UserError(f"У файлі вже є блок підсумків (рядок {r}). Візьміть файл після першої програми.")

    data_last = (total_row - 1) if total_row else last
    cats_found = [sh.get(r, CAT_COL) for r in range(head + 1, data_last + 1)]
    if not any(isinstance(v, (int, float)) for v in cats_found):
        raise UserError("У колонці C немає номерів категорій. Спочатку запустіть першу програму (Excel_Calc).")

    start = max((total_row + GAP_AFTER_TOTAL) if total_row else last + GAP_AFTER_TOTAL, last + 2)
    off = start - BLOCK_ROW0
    first_data, last_data = head + 1, start - 1               # діапазон для SUMIF (разом з рядками під «Разом»)
    L = lambda col: LET[col - 1]
    rng = lambda col: f"${L(col)}${first_data}:${L(col)}${last_data}"
    row = lambda dr: start + dr

    # ── що саме лежить у комірках блоку
    content: dict[tuple[int, int], object] = {}
    month = report_month(sh)
    if month:
        content[(0, 1)] = month
    content[(0, 5)], content[(0, 6)], content[(0, 9)] = LABEL_HEADER, "Кількість", "Сума"
    content[(0, 7)] = "формула = кільк./1000/2*0,175 (т.)"
    for dr, name, unit, code, cat in CAT_ROWS:
        r = row(dr)
        content[(dr, 1)], content[(dr, 3)], content[(dr, 4)], content[(dr, 5)] = name, unit, code, cat
        if cat != 6:                                           # «послуги» — кількості немає
            content[(dr, 6)] = ("f", f"SUMIF({rng(CAT_COL)},E{r},{rng(TARGET_COL)})")
        content[(dr, 9)] = ("f", f"SUMIF({rng(CAT_COL)},E{r},{rng(money_col)})")
        if unit == "грн.":
            content[(dr, 7)] = ("f", f"I{r}/1000")
        elif unit == "т":
            content[(dr, 7)] = ("f", f"F{r}/1000/2*0.175")
        else:
            content[(dr, 7)] = ("f", f"F{r}/2*0.175")
    r3 = row(3)                                                # рядок «Продукція інша» (3): додаткова вага
    content[(3, 10)] = ("f", f"F{r3}/2000*0.175")
    content[(3, 11)] = ("f", f"ROUND(J{r3},0)")
    tot = row(17)
    content[(17, 6)] = ("f", f"SUM(F{row(1)}:F{row(15)})")
    content[(17, 7)] = ("f", f"SUM(G{row(1)}:G{row(16)})")
    content[(17, 9)] = ("f", f"SUM(I{row(1)}:I{row(16)})")
    content[(19, 6)] = "1-ПЕ"
    content[(21, 6)], content[(21, 7)], content[(21, 8)] = 701, 702, 703

    # ── запис
    style_of = {(dr, c + 1): sid for dr, c, sid in CELLS}      # CELLS: колонки з 0, content: з 1
    default_sid = None
    for key in sorted(set(style_of) | set(content)):
        dr, c = key
        sid = style_of.get(key)
        if sid is None:                                        # значення без видимого стилю: беремо найпростіший
            if default_sid is None:
                default_sid = next(i for i, s in enumerate(STYLES) if not s["fill"] and not any(b[0] for b in s["bord"])
                                   and s["font"][3] == 0 and s["fmt"] == "General")
            sid = default_sid
        sh.put_block(row(dr), c, content.get(key), sid)
    for dr, h in HEIGHTS.items():
        sh.row_height(row(int(dr)), h)
    for dr, c1, c2 in MERGES:
        sh.merge(row(dr), c1 + 1, c2 + 1)
    for c, w in WIDTHS.items():
        sh.col_width(int(c) + 1, w)

    out_name = src.name[len("пораховано_"):] if src.name.startswith("пораховано_") else src.name
    out = src.with_name(OUT_PREFIX + out_name)
    if out.exists():
        out = src.with_name(f"{OUT_PREFIX}{Path(out_name).stem}_{dt.datetime.now():%Y%m%d_%H%M%S}{src.suffix}")
    try:
        sh.save(out)
    except PermissionError:
        raise UserError("Не вдалось зберегти результат. Закрийте файл у Excel і спробуйте ще раз.")

    # ── очікувані значення (рахуємо незалежно, для звірки та підсумку)
    exp = {}
    unaccounted = []
    no_qty = 0                                                 # рядки, де D = ######## (кількість не визначена)
    valid = {cat for _, _, _, _, cat in CAT_ROWS}
    tot_d = tot_g = Decimal(0)
    for r in range(first_data, data_last + 1):
        d, g, c = num(sh.get(r, TARGET_COL)), num(sh.get(r, money_col)), sh.get(r, CAT_COL)
        nm = sh.get(r, name_col)
        if isinstance(nm, str) and sh.get(r, TARGET_COL) == MARK:
            no_qty += 1
        is_item = isinstance(nm, str) and d is not None
        if is_item:
            tot_d += d
            tot_g += g or 0
        if isinstance(c, (int, float)) and int(c) == c and int(c) in valid:
            e = exp.setdefault(int(c), [Decimal(0), Decimal(0)])
            e[0] += d or 0
            e[1] += g or 0
        elif is_item:
            unaccounted.append((r, nm, c, d, g or Decimal(0)))
    extra_money = Decimal(0)                                   # ручні рядки між «Разом» і блоком (напр. послуги)
    for r in range((total_row or data_last) + 1, last_data + 1):
        c, g = sh.get(r, CAT_COL), num(sh.get(r, money_col))
        if isinstance(c, (int, float)) and int(c) in valid and g is not None:
            exp.setdefault(int(c), [Decimal(0), Decimal(0)])[1] += g
            extra_money += g
    return out, dict(start=start, month=month, exp=exp, unaccounted=unaccounted, tot_d=tot_d, tot_g=tot_g,
                     sum_f=sum(v[0] for v in exp.values()), sum_i=sum(v[1] for v in exp.values()),
                     extra_money=extra_money, no_qty=no_qty, fmt=src.suffix.lower())


# ───────────────────────── інтерфейс ─────────────────────────
def summary(out: Path, res: dict) -> str:
    f = lambda x: (f"{int(x):,}" if x == x.to_integral_value() else f"{x:,.2f}").replace(",", " ")
    lines = [f"Збережено: {out.name}", f"Блок додано з рядка {res['start']}"]
    if not res["month"]:
        lines.append("⚠ Не знайшов місяць у заголовку звіту — дату в A не заповнено.")
    lines.append(f"Кількість по категоріях разом: {f(res['sum_f'])} (по колонці D: {f(res['tot_d'])})")
    lines.append(f"Сума по категоріях разом: {f(res['sum_i'])} грн")
    if res["no_qty"]:
        lines.append(f"⚠ Рядків без розрахованої кількості ({MARK} у D): {res['no_qty']} — у F вони не входять.")
    if res["unaccounted"]:
        d = sum(x[3] for x in res["unaccounted"])
        g = sum(x[4] for x in res["unaccounted"])
        lines.append(f"\n⚠ Не потрапило в жодну категорію: {len(res['unaccounted'])} рядків "
                     f"({f(d)} шт, {f(g)} грн)")
        for r, n, c, dd, gg in res["unaccounted"][:12]:
            lines.append(f"  рядок {r}: {n} (C = {c})")
        if len(res["unaccounted"]) > 12:
            lines.append(f"  … і ще {len(res['unaccounted']) - 12}")
        lines.append("Додайте слова в reference.txt і запустіть обидві програми заново.")
    else:
        lines.append("Усі рядки розподілено по категоріях.")
    if res["fmt"] == ".xlsx":
        lines.append("\nФормули в .xlsx порахуються після відкриття в Excel.")
    return "\n".join(lines)


def main() -> None:
    for src in pick_files():
        try:
            out, res = process(src)
        except UserError as e:
            say("Не вийшло", f"{src.name}\n\n{e}", error=True)
            continue
        except Exception as e:
            say("Непередбачена помилка", f"{src.name}\n\n{type(e).__name__}: {e}", error=True)
            continue
        say("Готово", summary(out, res))
        if hasattr(os, "startfile"):
            try:
                os.startfile(out)
            except OSError:
                pass


if __name__ == "__main__":
    main()
