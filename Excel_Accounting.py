# -*- coding: utf-8 -*-
"""
Розрахунок кількості штук з оборотно-сальдової відомості (1С).

  «Візитка, 1000 шт»  +  оборот 62   ->   у колонку D записується 62000

Оформлення таке саме, як у старому скрипті:
  * результат — лише в колонці D, без додаткових колонок і аркушів;
  * загальна сума всіх порахованих рядків — у колонці D рядка «Разом»;
  * номер категорії товару — у колонці C (довідник: reference.txt або вбудований; 12, 13, 14 — жовті);
  * якщо рядок порахувати не можна — у D ставиться "########";
  * оригінал не змінюється, поруч зберігається копія «пораховано_<назва>».

Запуск: перетягнути файл .xls / .xlsx на програму (Excel_Calc.exe або Drop_file_here.bat);
        або двічі клікнути — відкриється вікно вибору файлу.
Рядки без «шт» у назві («Блокнот А6», «Етикетка 028») рахуються як 1 шт: результат = оборот.

Потрібні пакети:  pip install openpyxl xlrd xlutils xlwt
"""
from __future__ import annotations

import datetime as dt
import os
import re
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

TARGET_COL = 4                 # колонка D, як у старому скрипті
CAT_COL = 3                    # колонка C — номер категорії
YELLOW_CATEGORIES = {12, 13, 14}   # ці категорії виділяються жовтим
MARK = "########"              # позначка «порахувати не вдалось», як у старому скрипті
PREFIX = "пораховано_"
TOTAL_WORDS = {"разом", "всього", "усього", "итого", "итог"}

# число перед «шт / шт. / штук»; не береться, якщо воно частина іншого слова/числа
RUN_RE = re.compile(r"(?<![\w.,/\-])(\d+)\s*(?:шт|штук)\b\.?", re.IGNORECASE)


# Довідник категорій. Якщо поруч із програмою лежить файл reference.txt — використовується він
# (його можна редагувати, щоб додати нові слова), інакше — ця вбудована копія.
DEFAULT_REFERENCE = """1 - Банер,  плакат,  оголошення,  постер,  афіша
2 - календар, сітка квартальна
3 - візитка,  квиток,  табличка,  фірмовий бланк,  хенгер,  паспорт,  тейбл тент,  прапор,  бейдж,  картина,  сети для ресторанів,  вкладка,  інструкція,  картина, вкладиш
4 - Друковані аркуші,  листівка,  флаєр,  блок для записів,  меню,  картки,  прайс,  обкладинка,  сертифікат,  пам'ятка
5 - брошура,  буклет,  журнал,  посібник
6 - послуги
7 - коробка,  пакет,  упаковка
8 - конверти
9 - папка
10 - наклейка, наліпка, стікер
11 - бірка,  етикетка,  ярлик,  цінник,  воблер,  шелфтокер

12 - стакани
13 - каталог
14 - блокноти,  щоденник,  записні книжки,  записники
"""


class UserError(Exception):
    """Помилка, яку показуємо користувачу простими словами."""


# ───────────────────────── читання / запис (.xls і .xlsx) ─────────────────────────
class Sheet:
    """Єдиний інтерфейс: get(r, c) -> значення (рядки й колонки з 1), put(r, c, value), save(path)."""
    nrows: int

    def get(self, r: int, c: int): ...
    def put(self, r: int, c: int, value, center: bool = False, fill: str | None = None) -> None: ...
    def save(self, path: Path) -> None: ...


class XlsxSheet(Sheet):
    def __init__(self, path: Path):
        from openpyxl import load_workbook
        self.wb = load_workbook(path)
        self.ws = self.wb.active
        self.nrows = self.ws.max_row

    def get(self, r, c):
        return self.ws.cell(r, c).value

    def put(self, r, c, value, center=False, fill=None):
        from copy import copy
        from openpyxl.styles import PatternFill
        cell = self.ws.cell(r, c)
        if type(cell).__name__ == "MergedCell":
            raise UserError(f"Комірка {cell.coordinate} входить в об'єднану область — не можу записати результат.")
        cell.value = value
        if center:
            al = copy(cell.alignment)
            al.horizontal = "center"
            cell.alignment = al
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)

    def save(self, path):
        self.wb.save(path)


class XlsSheet(Sheet):
    def __init__(self, path: Path):
        import xlrd
        from xlutils.filter import XLRDReader, XLWTWriter, process
        self.rb = xlrd.open_workbook(str(path), formatting_info=True)
        self.rs = self.rb.sheet_by_index(0)
        writer = XLWTWriter()
        process(XLRDReader(self.rb, "unknown.xls"), writer)
        self.wb = writer.output[0][1]
        self.ws = self.wb.get_sheet(0)
        self.ws.header_str = b""                    # xlwt за замовчуванням додає колонтитули для друку
        self.ws.footer_str = b""
        # xlutils не переносить власну палітру кольорів файлу — без цього кольори «пливуть»
        for idx, rgb in enumerate(self.rb.palette_record or [], 8):
            if idx <= 63:
                self.wb.set_colour_RGB(idx, *rgb)
        self.style_list = writer.style_list
        self.nrows = self.rs.nrows
        self._merged = set()
        for rlo, rhi, clo, chi in self.rs.merged_cells:
            for rr in range(rlo, rhi):
                for cc in range(clo, chi):
                    self._merged.add((rr, cc))
                    # xlwt заливає всю об'єднану область стилем першої комірки — повертаємо оригінальні рамки
                    xf = self.rs.cell_xf_index(rr, cc)
                    row = self.ws._Worksheet__rows.get(rr)
                    cell = row._Row__cells.get(cc) if row is not None else None
                    if cell is not None and hasattr(cell, "xf_idx"):
                        cell.xf_idx = self.wb.add_style(writer.style_list[xf])

    def get(self, r, c):
        if r > self.rs.nrows or c > self.rs.ncols:
            return None
        v = self.rs.cell_value(r - 1, c - 1)
        return None if v == "" else v

    def put(self, r, c, value, center=False, fill=None):
        if (r - 1, c - 1) in self._merged:
            raise UserError(f"Комірка в рядку {r} входить в об'єднану область — не можу записати результат.")
        row = self.ws._Worksheet__rows.get(r - 1)
        old = row._Row__cells.get(c - 1) if row is not None else None
        xf = getattr(old, "xf_idx", None)           # зберігаємо існуючий формат комірки
        if center or fill:                          # той самий стиль + вирівнювання / заливка
            import copy
            import xlwt
            try:
                base = self.style_list[self.rs.cell_xf_index(r - 1, c - 1)]
            except IndexError:
                base = xlwt.XFStyle()
            st = copy.deepcopy(base)
            if center:
                st.alignment.horz = xlwt.Alignment.HORZ_CENTER
            if fill:                                # заливку задаємо стандартним жовтим (індекс 5)
                st.pattern.pattern = xlwt.Pattern.SOLID_PATTERN
                st.pattern.pattern_fore_colour = 5
            xf = self.wb.add_style(st)
        self.ws.write(r - 1, c - 1, value)
        if xf is not None:
            self.ws._Worksheet__rows[r - 1]._Row__cells[c - 1].xf_idx = xf

    def save(self, path):
        self.wb.save(str(path))


def open_sheet(path: Path) -> Sheet:
    ext = path.suffix.lower()
    try:
        if ext in (".xlsx", ".xlsm"):
            return XlsxSheet(path)
        if ext == ".xls":
            return XlsSheet(path)
    except ImportError as e:
        raise UserError(f"Не вистачає пакета: {e.name}.\nВстановіть:  pip install openpyxl xlrd xlutils xlwt")
    raise UserError(f"Формат «{ext}» не підтримується. Потрібен .xls або .xlsx.")


# ───────────────────────── допоміжні функції ─────────────────────────
def to_decimal(v) -> Decimal | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).strip()
    for ch in (" ", "\u00a0", "\u202f"):
        s = s.replace(ch, "")
    try:
        return Decimal(s.replace(",", "."))
    except InvalidOperation:
        return None


SHT_WORD = re.compile(r"(?<!\w)(?:шт|штук)(?!\w)", re.IGNORECASE)
END_RUN = re.compile(r"(\d+)\s*(?:шт|штук)\.?\s*$", re.IGNORECASE)


def parse_qty(name: str) -> tuple[int | None, str]:
    """(штук на одиницю, вид): 'run' — «… 1000 шт»; 'single' — «шт» у назві немає, це 1 шт;
    'bad' — «шт» є, але тираж неоднозначний (рахувати не будемо)."""
    if not SHT_WORD.search(name):
        return 1, "single"
    if len(SHT_WORD.findall(name)) == 1:
        m = END_RUN.search(name)
        # «1 000 шт», «А4 300 шт» — число стоїть поруч з іншим числом: неоднозначно
        if m and not re.search(r"\d[\s\u00a0\u202f]*$", name[: m.start()]) and int(m.group(1)) > 0:
            return int(m.group(1)), "run"
    return None, "bad"


# ───────────────────────── категорії (колонка C) ─────────────────────────
_END = ("ими", "ами", "ями", "ого", "ому", "ою", "ої", "ий", "ій", "им", "их", "ім", "ів", "ей", "ам", "ям", "ах", "ях")


def tokenize(text: str) -> list[str]:
    text = re.sub(r"[’'ʼ`´]", "", text.casefold())      # пам'ятка = памятка
    return re.findall(r"[a-zа-яіїєґ0-9]+", text)


def stems(word: str) -> frozenset[str]:
    """Основа слова, щоб однина/множина й відмінки збігались: візитка/візитки/візиток, квиток/квитки."""
    if not re.search(r"[а-яіїєґ]", word):
        return frozenset({word})
    s = word
    for e in _END:
        if s.endswith(e) and len(s) - len(e) >= 3:
            s = s[: -len(e)]
            break
    else:
        t = re.sub(r"[аеєиіїоюяьу]+$", "", s)
        s = t if len(t) >= 2 else s
    out = {s}
    if s.endswith("ок") and len(s) >= 5:                # випадний голосний: квиток -> квитк
        out.add(s[:-2] + "к")
    return frozenset(out)


def parse_reference(text: str) -> list[tuple[int, tuple[frozenset, ...]]]:
    """Рядки виду «3 - візитка,  квиток, ...» -> [(3, (стеми слова,)), ...]. Фрази з кількох слів теж підтримані."""
    refs = []
    for line in text.splitlines():
        m = re.match(r"^\s*(\d+)\s*[-–—:]\s*(.+?)\s*$", line)
        if not m:
            continue
        for item in m.group(2).split(","):
            toks = tokenize(item)
            if toks:
                refs.append((int(m.group(1)), tuple(stems(t) for t in toks)))
    return refs


def app_dir() -> Path:
    return Path(sys.executable).parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent


def load_reference() -> tuple[list, str]:
    text, source = DEFAULT_REFERENCE, "вбудований довідник"
    path = app_dir() / "reference.txt"
    if path.exists():
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text, source = path.read_text(encoding=enc), "reference.txt"
                break
            except UnicodeDecodeError:
                continue
    refs = parse_reference(text)
    if not refs:
        raise UserError(f"У довіднику категорій ({source}) немає жодного рядка виду «3 - візитка, квиток».")
    return refs, source


def _hits(tokens: list[frozenset], refs) -> dict[int, int]:
    hits: dict[int, int] = {}
    for cat, phrase in refs:
        k = len(phrase)
        for i in range(len(tokens) - k + 1):
            if all(tokens[i + j] & phrase[j] for j in range(k)):
                hits.setdefault(cat, i)
                break
    return hits


def find_category(name: str, refs) -> tuple[int | None, str]:
    """(номер категорії, 'ok') | (None, 'unknown') | (None, 'ambiguous').
    Спочатку дивимось на частину назви до першої коми («Візитка, 250 шт» -> «Візитка»)."""
    head, _, _ = name.partition(",")
    for part in (head, name):
        hits = _hits([stems(t) for t in tokenize(part)], refs)
        if len(hits) == 1:
            return next(iter(hits)), "ok"
        if len(hits) > 1:
            return None, "ambiguous"
    return None, "unknown"


def find_columns(sh: Sheet) -> tuple[int, int, int]:
    """(колонка назв, колонка оборотів-Дебет, рядок заголовка) — за заголовками, а не за номерами."""
    for r in range(1, min(sh.nrows, 40) + 1):
        for c in range(1, 15):
            v = sh.get(r, c)
            if not (isinstance(v, str) and v.strip().casefold().startswith("обороти за період")):
                continue
            for rr in range(r + 1, r + 4):
                for cc in (c, c + 1):
                    d = sh.get(rr, cc)
                    if isinstance(d, str) and d.strip().casefold() == "дебет":
                        name_col = 1
                        for x in range(1, 15):
                            h = sh.get(r, x)
                            if isinstance(h, str) and h.strip().casefold() == "рахунок":
                                name_col = x
                                break
                        return name_col, cc, rr
    raise UserError("Не знайшов у файлі заголовок «Обороти за період» → «Дебет».\n"
                    "Схоже, це інший звіт. Скрипт не вгадує колонки, щоб не порахувати неправильно.")


# ───────────────────────── основна обробка ─────────────────────────
def process(src: Path) -> tuple[Path, dict]:
    sh = open_sheet(src)
    name_col, cnt_col, head = find_columns(sh)

    # рядок «Разом»
    total_row = total_cnt = None
    for r in range(head + 1, sh.nrows + 1):
        v = sh.get(r, name_col)
        if isinstance(v, str) and v.strip().casefold().strip(" :.") in TOTAL_WORDS:
            total_row, total_cnt = r, to_decimal(sh.get(r, cnt_col))
            break
    last = (total_row - 1) if total_row else sh.nrows

    # захист: колонки C і D у рядках даних мають бути порожніми
    for r in range(head + 1, (total_row or last) + 1):
        for col, letter in ((CAT_COL, "C"), (TARGET_COL, "D")):
            if sh.get(r, col) not in (None, ""):
                raise UserError(f"Колонка {letter} уже містить дані (рядок {r}). Щоб нічого не затерти, "
                                f"скрипт зупинено.")
    refs, ref_source = load_reference()

    results: dict[int, object] = {}
    cats: dict[int, object] = {}                    # рядок -> номер категорії або MARK
    cat_bad: list[tuple[int, str, str]] = []        # рядки без однозначної категорії
    bad: list[tuple[int, str, object]] = []         # що не вдалось порахувати
    sum_cnt = Decimal(0)                            # для звірки з «Разом»
    sum_pieces = Decimal(0)
    n_ok = 0
    n_single = 0                                    # рядки без «шт» у назві (рахуються як 1 шт)

    for r in range(head + 1, last + 1):
        raw_name = sh.get(r, name_col)
        if raw_name is None or isinstance(raw_name, (int, float)):
            continue                                # порожні рядки, рахунок «26»
        name = " ".join(str(raw_name).split())
        raw_cnt = sh.get(r, cnt_col)
        cnt = to_decimal(raw_cnt)
        qty, kind = parse_qty(name)

        if cnt is not None and total_cnt is not None and cnt == total_cnt and kind == "single":
            continue                                # проміжні підсумки («Обл. НДС»)
        if raw_cnt is None and kind == "single":
            continue                                # підписи на кшталт «Номенклатура»

        cat, status = find_category(name, refs)
        cats[r] = cat if cat is not None else MARK
        if cat is None:
            cat_bad.append((r, name, status))

        if cnt is not None:
            sum_cnt += cnt
        if (qty is not None and cnt is not None
                and cnt > 0 and cnt == cnt.to_integral_value()):
            value = qty * int(cnt)
            results[r] = value
            sum_pieces += value
            n_ok += 1
            n_single += kind == "single"
        else:
            results[r] = MARK
            bad.append((r, name, raw_cnt))

    if total_row:                                   # загальна сума — у колонку D рядка «Разом»
        results[total_row] = int(sum_pieces)
    for r, v in results.items():
        sh.put(r, TARGET_COL, v)
    for r, v in cats.items():                       # категорії 12, 13, 14 — жовта комірка
        sh.put(r, CAT_COL, v, center=True, fill="FFFF00" if v in YELLOW_CATEGORIES else None)

    out = src.with_name(f"{PREFIX}{src.stem}{src.suffix}")
    if out.exists():
        out = src.with_name(f"{PREFIX}{src.stem}_{dt.datetime.now():%Y%m%d_%H%M%S}{src.suffix}")
    try:
        sh.save(out)
    except PermissionError:
        raise UserError("Не вдалось зберегти результат. Закрийте файл у Excel і спробуйте ще раз.")

    if total_cnt is None:
        recon = None
    else:
        recon = (sum_cnt == total_cnt, sum_cnt, total_cnt)
    return out, dict(cat_bad=cat_bad, ref_source=ref_source, n_cat=len(cats) - len(cat_bad), total_written=bool(total_row), ok=n_ok, single=n_single, bad=bad, pieces=sum_pieces, recon=recon)


# ───────────────────────── інтерфейс ─────────────────────────
def _tk():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        return root, filedialog, messagebox
    except Exception:
        return None, None, None


def pick_files() -> list[Path]:
    if len(sys.argv) > 1:
        return [Path(a) for a in sys.argv[1:]]
    root, fd, _ = _tk()
    if fd:
        names = fd.askopenfilenames(title="Оберіть відомість (.xls / .xlsx)",
                                    filetypes=[("Excel", "*.xls *.xlsx *.xlsm")])
        root.destroy()
        return [Path(n) for n in names]
    p = input("Шлях до файлу: ").strip().strip('"')
    return [Path(p)] if p else []


def say(title: str, text: str, error: bool = False) -> None:
    root, _, mb = _tk()
    if mb:
        (mb.showerror if error else mb.showinfo)(title, text)
        root.destroy()
    else:
        print(f"\n== {title} ==\n{text}")


def summary(out: Path, res: dict) -> str:
    fmt = lambda x: f"{x:,.0f}".replace(",", " ")
    lines = [f"Збережено: {out.name}",
             f"Пораховано рядків: {res['ok']}   (разом {fmt(res['pieces'])} шт)",
             f"З них без «шт» у назві (оборот = кількість штук): {res['single']}"]
    if not res["total_written"]:
        lines.append("⚠ Рядок «Разом» не знайдено — загальну суму в файл не записано.")
    if res["recon"]:
        ok, s, t = res["recon"]
        lines.append("Звірка з «Разом»: ОК" if ok else
                     f"⚠ Звірка з «Разом» НЕ збіглась: по рядках {s}, у «Разом» {t}")
    if res["bad"]:
        lines.append(f"\nПозначено «{MARK}» (у загальну суму НЕ входять, треба глянути вручну): {len(res['bad'])}")
        for r, n, c in res["bad"][:15]:
            lines.append(f"  рядок {r}: {n} (оборот {c})")
        if len(res["bad"]) > 15:
            lines.append(f"  … і ще {len(res['bad']) - 15}")
    lines.append(f"\nКатегорії (колонка C): визначено {res['n_cat']}, довідник: {res['ref_source']}")
    if res["cat_bad"]:
        lines.append(f"Категорію НЕ визначено ({MARK} у колонці C): {len(res['cat_bad'])}")
        why = {"unknown": "немає в довіднику", "ambiguous": "підходить кілька категорій"}
        for r, n, st in res["cat_bad"][:15]:
            lines.append(f"  рядок {r}: {n} — {why[st]}")
        if len(res["cat_bad"]) > 15:
            lines.append(f"  … і ще {len(res['cat_bad']) - 15}")
        lines.append("Щоб додати слово, впишіть його в потрібний рядок файлу reference.txt (поруч з програмою).")
    return "\n".join(lines)


def main() -> None:
    for src in pick_files():
        try:
            out, res = process(src)
        except UserError as e:
            say("Не вийшло", f"{src.name}\n\n{e}", error=True)
            continue
        except Exception as e:
            say("Непередбачена помилка", f"{src.name}\n\n{type(e).__name__}: {e}", error=True)
            continue
        say("Готово", summary(out, res))
        if hasattr(os, "startfile"):                # Windows: одразу відкрити результат
            try:
                os.startfile(out)
            except OSError:
                pass


if __name__ == "__main__":
    main()
